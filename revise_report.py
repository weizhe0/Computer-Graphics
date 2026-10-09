from pathlib import Path
import shutil
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\taiwz\Desktop\GHL\tcg_project.docx")
OUTPUT = Path(r"C:\Users\taiwz\Desktop\GHL\tcg_project_revised_highlighted.docx")


def set_highlighted_paragraph(paragraph, text):
    alignment = paragraph.alignment
    style = paragraph.style
    paragraph.clear()
    paragraph.style = style
    paragraph.alignment = alignment
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def find_paragraph(doc, exact=None, startswith=None):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if exact is not None and text == exact:
            return paragraph
        if startswith is not None and text.startswith(startswith):
            return paragraph
    raise ValueError(f"Paragraph not found: exact={exact!r}, startswith={startswith!r}")


def replace_paragraph(doc, text, replacement, exact=True):
    paragraph = find_paragraph(doc, exact=text if exact else None,
                               startswith=None if exact else text)
    return set_highlighted_paragraph(paragraph, replacement)


def insert_before(target, text, style=None):
    new_p = OxmlElement("w:p")
    target._p.addprevious(new_p)
    paragraph = Paragraph(new_p, target._parent)
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def set_cell(cell, text):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


doc = Document(SOURCE)

# Front matter and project-development record.
replace_paragraph(doc, "GANTT CHART (SAMPLE)", "PROJECT DEVELOPMENT TIMELINE")

timeline = [
    ("Week 3", "Concept planning and task allocation",
     "All members: selected the Minecraft blind-box battle concept and divided character, environment, gameplay, documentation, and presentation responsibilities."),
    ("Week 4", "Character modelling",
     "Tam Yu Siong: Creeper modelling and texture research. Elysa Lee Xing Wan: Enderman modelling and visual references. Tai Wei Zhe: OpenGL class structure and hierarchical transformations. Chong Jia Ming: Blaze modelling and arena asset planning."),
    ("Week 5", "Skill design and battle mechanics",
     "All members: designed two personality-based skills for each character and reviewed attack distance, damage, cooldown, and control balance."),
    ("Weeks 6-7", "Animation and rendering effects",
     "Tai Wei Zhe: integrated movement, skill logic, damage feedback, knockback, and camera behaviour. Other members: refined character proportions, procedural textures, particle appearance, and environment presentation."),
    ("Weeks 8-9", "Dynamic environment and user interface",
     "Team: implemented the character-selection menu, health and cooldown HUD, reactive lighting, sky colour changes, arena effects, skill warnings, and end-state animations."),
    ("Week 10", "Testing and refinement",
     "All members: tested every character pairing, corrected controls, improved walking and casting animations, balanced skill range and damage, and verified window/full-screen presentation."),
    ("Weeks 11-12", "Report, video, and demonstration",
     "Tam Yu Siong and Elysa Lee Xing Wan: report content and screenshots. Tai Wei Zhe: source-code verification and technical documentation. Chong Jia Ming: video editing and demonstration preparation. All members: final review and rehearsal."),
]
table = doc.tables[0]
for row_index, values in enumerate(timeline, start=1):
    for col_index, value in enumerate(values):
        set_cell(table.rows[row_index].cells[col_index], value)

# Chapter 1: replace repeated or inaccurate prose.
replace_paragraph(
    doc,
    "Blind Boxes Collection: Minecraft Battle is a real-time interactive application",
    "Blind Boxes Collection: Minecraft Battle is a real-time local two-player application developed for TCG 6223 Computer Graphics. Inspired by Minecraft's hostile mobs, the project presents Creeper, Enderman, and Blaze as selectable blind-box characters in a three-dimensional battle arena. The program is written in C++ using the OpenGL fixed-function pipeline and GLUT for windowing and input.",
    exact=False,
)
replace_paragraph(
    doc,
    "The project is based on the collectible blind box concept:",
    "The project combines the surprise and collectability of a blind-box series with an interactive battle game. Players select two different characters from the main menu, enter the arena, move in real time, and use character-specific close-range or long-range skills. This concept allows the project to demonstrate both object modelling and responsive computer-graphics animation.",
    exact=False,
)
replace_paragraph(
    doc,
    "The uniqueness of this project is that all the virtual objects are created procedurally.",
    "All virtual objects are created programmatically without external 3D model files or code generators. Character parts are assembled from custom box primitives whose vertices, faces, normals, and texture coordinates are coded manually. Character, terrain, and blind-box textures are generated at run time and uploaded with glGenTextures(), glBindTexture(), and glTexImage2D(). This approach satisfies the project requirement while demonstrating direct control over geometry, transformations, materials, and texture mapping.",
    exact=False,
)
replace_paragraph(
    doc,
    "The three playable characters are modelled after hostile mobs",
    "The three playable characters are inspired by hostile mobs from Minecraft. Each model uses hierarchical transformations with glPushMatrix() and glPopMatrix(), allowing heads, limbs, and accessories to move relative to the character body. This structure supports walking, casting, damage, knockback, death, and victory animations.",
    exact=False,
)
replace_paragraph(
    doc,
    "It is a four-legged animal that with a high, thin body",
    "The Creeper has a tall, narrow body, four short legs, a mottled green procedural texture, and a manually constructed pixel face. Its primary skill, Explosion, includes a 0.28-second warning phase in which the body expands and flashes green-white before releasing a radial shockwave, particles, knockback, camera shake, and a scorched crater. Its secondary skill, Powder Cloud, creates a lingering green hazard that applies periodic poison damage while the opponent remains inside it.",
    exact=False,
)
replace_paragraph(
    doc,
    "The Enderman has the tall, slender, long-limbed body",
    "The Enderman has a tall, slender body, long articulated limbs, a dark procedural texture, and emissive purple eyes produced with GL_EMISSION. Teleport Slash moves the Enderman behind a nearby opponent and applies a close-range strike, while Void Beam provides a ranged attack. During casting, both arms rise and surrounding purple particles spiral inward, strengthening the visual connection between the character pose and the released skill.",
    exact=False,
)
replace_paragraph(
    doc,
    "The Blaze is a floating, fire-wreathed being",
    "The Blaze is a floating fire-based character with a segmented core, a pixel-patterned head, and eight animated rods arranged at different heights. Fire Burst is a close-range radial attack. Fireball is a long-range projectile that first gathers flame in front of the Blaze for 0.24 seconds, then follows an arched trajectory toward a visible warning circle. Its small explosion area rewards accurate timing and allows the opponent to dodge.",
    exact=False,
)
replace_paragraph(
    doc,
    "A two-player battle mode is available when playing locally.",
    "The application provides a local two-player battle mode. Player 1 moves with W, A, S, and D and activates skills with G and H. Player 2 moves with the arrow keys and activates skills with 8 and 9. Before battle, both players select different characters from a keyboard-controlled main menu. Health bars, cooldown indicators, range limits, distance-based damage, poison ticks, and parabolic knockback provide readable combat feedback. After a winner is determined, ENTER starts a new round.",
    exact=False,
)
replace_paragraph(
    doc,
    "The battle is played out on a procedurally textured arena floor.",
    "The battle takes place in a procedurally textured arena surrounded by fences, block hills, stones, blind boxes, and a colour-changing skybox. The environment reacts to combat: Creeper produces green-white flashes, grey-green smoke, cracks, debris, and a scorched crater; Enderman darkens the arena with purple pulses; Blaze warms the sky and leaves a burning area. When either player has low health, the lighting gradually darkens. These transitions fade smoothly rather than changing abruptly.",
    exact=False,
)
replace_paragraph(
    doc,
    "Custom Primitive Rendering",
    "Custom Primitive Rendering - PrimitiveRenderer::drawBox() defines reusable box geometry with manually specified vertices, face indices, normals, and texture coordinates instead of relying on GLUT solid-object functions.",
    exact=False,
)

# Chapter 2 corrections.
replace_paragraph(
    doc,
    "[Insert Image: Arena and Skybox overview]",
    "Figure 2.1: Arena and skybox overview showing the procedural grass floor, boundary fence, block hills, stones, blind boxes, and dynamic sky.",
)
replace_paragraph(
    doc,
    "An environment is composed of a skybox",
    "The environment consists of a skybox, a textured ground plane, wooden fences, layered block hills, stones, arena markings, and blind-box decorations. The skybox is assembled from GL_QUADS with a size of 28.0 units. Its upper and lower colours are recalculated from the active Creeper, Enderman, Blaze, and low-health environment states. The ground extends from -ARENA_HALF_SIZE to +ARENA_HALF_SIZE, where ARENA_HALF_SIZE is 12.0 units. A 32 x 32 procedural grass texture is repeated across the floor, and translucent lines and markings improve depth and spatial awareness.",
    exact=False,
)
replace_paragraph(
    doc,
    "The Creeper's body is created using",
    "The Creeper body is drawn with PrimitiveRenderer::drawBox(0.95f, 1.45f, 0.72f) and translated upward by 1.38f plus walkBodyBob(). The head uses a 1.18f cube placed at y = 2.48f plus idle and walking offsets. During Explosion casting, the complete model scales outward and upward according to the charge value, while an emissive green-white material flashes at increasing intensity. The body and head use a 16 x 16 procedurally generated pixel texture.",
    exact=False,
)
replace_paragraph(
    doc,
    "The Blaze's core consists of three decreasing boxes",
    "The Blaze core is assembled from three boxes ranging from 0.44f x 1.22f x 0.44f to 0.34f x 0.34f x 0.34f. The model floats using idleFloat() = sin(mParticleTimer * 2.1f) * 0.08f. Its head is a 1.18f x 1.08f x 1.18f box at y = 2.34f. The face and side panels are generated from 8 x 8 colour-code arrays; each code selects a gold, amber, brown, red, or black material for a small drawFlatRect() tile.",
    exact=False,
)
replace_paragraph(
    doc,
    "The Blaze has 8 rods orbiting around it",
    "The Blaze has eight rods positioned around its core. Each rod is built by drawSegmentedRod(), which draws eight smaller boxes with alternating yellow, gold, and orange materials. drawBlazeRod() applies translation, yaw, roll, and sinusoidal sway to each rod. During the Blaze victory animation, the complete rod group rotates rapidly and pulses outward, making the celebration clearly different from its normal idle animation.",
    exact=False,
)

# Chapter 3: update implementation details to the final code.
replace_paragraph(
    doc,
    'The "first effect" shows the Creeper initiating its attack.',
    "During the 0.28-second activation phase, the Creeper body gradually expands and flashes green-white with increasing frequency. Explosion particles remain stationary during this warning phase, preventing the detonation effect from appearing too early. The animation uses delta time (dt), so its timing is consistent across different frame rates.",
    exact=False,
)
replace_paragraph(
    doc,
    "The second effect demonstrates the result of the explosion animation.",
    "At the impact point, a brief white core flash appears before a green shockwave expands to its maximum radius in approximately 0.20 seconds using cubic easing. A total of 220 particles are released. Bright and dark fragments use stronger gravity and bounce on the floor, while grey smoke uses lower gravity and drag so that it rises and disperses slowly. The same impact time triggers damage, distance-based knockback, camera shake, green-white environment lighting, cracks, debris, and a temporary scorched crater. Powder Cloud remains active as a translucent green poison zone that applies damage in timed ticks.",
    exact=False,
)
replace_paragraph(
    doc,
    'The "Teleport Slash" bypasses continuous translation.',
    "Teleport Slash first checks that the opponent is within the maximum teleport range. It calculates a destination behind the opponent from the opponent's facing vector, clamps the position to the arena boundary, and leaves a translucent after-image at the original location. During the cast, the Enderman's arms rise and purple particles spiral inward. A rapid visibility flicker and expanding rings reinforce the spatial transition before the close-range damage is resolved.",
    exact=False,
)
replace_paragraph(
    doc,
    'The "Void Beam" utilizes parametric equations',
    "Void Beam is a range-limited line attack drawn between the Enderman and the recorded target position. Multiple additive-blended lines, moving offsets, impact rings, and purple particles create the beam. The arm-raising pose and gathering particle animation are shared with the casting phase so that the character motion visibly prepares the attack.",
    exact=False,
)
replace_paragraph(
    doc,
    'The "Fireball" calculates a parabolic trajectory',
    "Fireball begins with a 0.24-second casting phase in which flame particles spiral toward a bright point in front of the Blaze. After launch, the projectile interpolates between its start and target positions. A half-sine vertical offset, arcHeight = sin(travelAge * PI) * 1.35f, creates the arched trajectory. A red-orange warning circle displays the 1.5-unit explosion radius so that the opponent can move away before impact.",
    exact=False,
)
replace_paragraph(
    doc,
    "Upon reaching the impact time",
    "At impact, the fireball changes from projectile rendering to concentric orange-red explosion rings, sparks, smoke, and a temporary burning ground area. Damage decreases from approximately 28 at the centre to approximately 7 near the edge, and knockback strength follows the same distance factor. Fire Burst remains an immediate close-range radial attack centred on the Blaze.",
    exact=False,
)
replace_paragraph(
    doc,
    "Rather than using static light values",
    "Battlefield updates the ambient and diffuse components of GL_LIGHT0, the skybox colours, and glClearColor() every frame. Creeper uses a dedicated three-stage response: green-white charging, a short white/yellow-green impact flash, and approximately three seconds of grey-green smoke darkening. Blaze uses an independent orange-red state connected to fire and burning areas. The values decay over time to restore the default environment smoothly.",
    exact=False,
)
replace_paragraph(
    doc,
    "Similarly, ethereal skills trigger alternative color pulses.",
    "Enderman skills trigger a purple-black environment state with expanding ground rings and flickering particles. This effect is independent of the Creeper and Blaze colour states. In addition, the Battlefield class monitors both health values; below 35 health, ambient and sky brightness gradually decrease to increase tension without hiding the characters or interface.",
    exact=False,
)

# Add final animation/feedback documentation before Chapter 4.
chapter4 = find_paragraph(doc, exact="CHAPTER 4: USER MANUAL INSTRUCTIONS")
insert_before(chapter4, "3.5 Combat Feedback and End-State Animation", "Heading 3")
insert_before(
    chapter4,
    "When a character takes damage, its materials temporarily shift toward bright red. Poison damage uses shorter repeated flashes, while airborne knockback keeps the target red until landing. Creeper and Fireball explosions apply horizontal velocity away from the impact point together with vertical velocity; gravity then produces a visible parabolic flight and landing instead of an instantaneous position change.",
)
insert_before(
    chapter4,
    "When health reaches zero, the defeated character smoothly falls sideways and becomes darker. The winner performs a continuous jumping and turning animation; Blaze additionally accelerates and expands its orbiting rods. The result overlay displays the winner and waits for ENTER, allowing the players to observe the final animation before starting a new round.",
)

# Update user instructions and controls.
replace_paragraph(
    doc,
    "In battle each player will have different controls",
    "During battle, Player 1 moves with W, A, S, and D and uses G and H for skills. Player 2 moves with the arrow keys and uses 8 and 9 for skills. The separate bindings reduce input conflict when two players share one keyboard. When a round ends, ENTER resets health, positions, cooldowns, environment effects, and victory/death states.",
    exact=False,
)
set_cell(doc.tables[2].rows[5].cells[2], "G")
set_cell(doc.tables[2].rows[6].cells[2], "H")
set_cell(doc.tables[2].rows[11].cells[2], "8")
set_cell(doc.tables[2].rows[12].cells[2], "9")
for row in (5, 6, 11, 12):
    set_cell(doc.tables[2].rows[row].cells[3],
             doc.tables[2].rows[row].cells[3].text.replace("Explosion Ring", "Explosion"))

# Add end-of-round ENTER instruction to the system table.
system_table = doc.tables[1]
new_row = system_table.add_row()
set_cell(new_row.cells[0], "End of Round")
set_cell(new_row.cells[1], "ENTER")
set_cell(new_row.cells[2], "Restart the battle after the winner and final animations are displayed.")

# Update screenshot text and add conclusion/references.
replace_paragraph(
    doc,
    "Figure 5.3:",
    "Figure 5.3: Skill-action screenshot showing a character casting animation, particles, hit feedback, dynamic lighting, and a reactive arena effect. The visual components are updated using frame-rate-independent timing.",
    exact=False,
)

video_heading = find_paragraph(doc, exact="VIDEO CLIP")
insert_before(video_heading, "CHAPTER 6: CONCLUSION", "Heading 2")
insert_before(
    video_heading,
    "Blind Boxes Collection: Minecraft Battle fulfils the project objective by combining original OpenGL object construction, hierarchical modelling, procedural texture generation, lighting, material effects, animation, and real-time interaction. Three distinct characters exceed the minimum requirement, and each character has two skills with different ranges, cooldowns, visual identities, and environmental responses. The final system demonstrates how graphics techniques can support both visual presentation and readable gameplay.",
)
insert_before(video_heading, "REFERENCES AND ACKNOWLEDGEMENT", "Heading 2")
insert_before(
    video_heading,
    "Minecraft character names and visual characteristics are acknowledged as inspiration from Mojang Studios and Microsoft. All OpenGL character geometry, procedural textures, battle logic, animations, particles, interface elements, and environment effects in this project were implemented by the student team without downloaded 3D models or generated OpenGL model code. Technical reference material included the OpenGL/GLU and GLUT function documentation used in the TCG 6223 laboratory exercises.",
)
replace_paragraph(
    doc,
    "Demonstration Link:",
    "Demonstration Link (video must be less than 5 minutes): [INSERT FINAL VIDEO URL HERE]",
    exact=False,
)

# Ask Word to refresh the real TOC field on opening.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.save(OUTPUT)

# Replace cover text stored inside text boxes and highlight the modified runs.
with tempfile.TemporaryDirectory() as tmp_dir:
    temp_path = Path(tmp_dir)
    with zipfile.ZipFile(OUTPUT, "r") as source_zip:
        source_zip.extractall(temp_path)

    document_xml = temp_path / "word" / "document.xml"
    from lxml import etree

    tree = etree.parse(str(document_xml))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for paragraph in tree.xpath("//w:p", namespaces=ns):
        text_nodes = paragraph.xpath(".//w:t", namespaces=ns)
        combined = "".join(node.text or "" for node in text_nodes)
        if "Project for Last Name" not in combined:
            continue

        replacement = combined.replace(
            "Project for Last Name",
            "Blind Boxes Collection: Minecraft Battle | Lab Section: [INSERT LAB SECTION]",
        )
        if text_nodes:
            text_nodes[0].text = replacement
            for node in text_nodes[1:]:
                node.text = ""

        for run in paragraph.xpath(".//w:r", namespaces=ns):
            rpr = run.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                run.insert(0, rpr)
            highlight = rpr.find(qn("w:highlight"))
            if highlight is None:
                highlight = OxmlElement("w:highlight")
                rpr.append(highlight)
            highlight.set(qn("w:val"), "yellow")
    tree.write(str(document_xml), xml_declaration=True, encoding="UTF-8", standalone="yes")

    rebuilt = OUTPUT.with_suffix(".tmp.docx")
    with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as output_zip:
        for path in temp_path.rglob("*"):
            if path.is_file():
                output_zip.write(path, path.relative_to(temp_path))
    shutil.move(rebuilt, OUTPUT)

print(OUTPUT)
